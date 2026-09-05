"""DOCX / plain text / Markdown / tabular file parsing."""
from __future__ import annotations

import csv
import io
import json

from ...core.errors import UnsupportedContent
from ...core.provenance import Passage, Source, SourceKind
from ...core.registry import Provider, ProviderStatus
from ...core.text import chunk, normalise

CAPABILITY = "document_parse"


class TextParser(Provider):
    capability = CAPABILITY
    name = "text"
    requires_credentials = False
    priority = 20
    extensions = (".txt", ".md", ".markdown", ".rst", ".log", ".json", ".xml", ".html", ".htm")

    def status(self) -> ProviderStatus:
        return ProviderStatus(True, details={"extensions": list(self.extensions)})

    def parse(self, data: bytes, filename: str) -> Source:
        try:
            text = data.decode("utf-8")
        except UnicodeDecodeError:
            text = data.decode("latin-1", errors="replace")
        if filename.lower().endswith((".html", ".htm")):
            from bs4 import BeautifulSoup

            soup = BeautifulSoup(text, "lxml")
            for tag in soup.find_all(["script", "style"]):
                tag.decompose()
            text = soup.get_text("\n")
        if filename.lower().endswith(".json"):
            try:
                text = json.dumps(json.loads(text), indent=2, ensure_ascii=False)
            except json.JSONDecodeError:
                pass
        text = normalise(text)
        src = Source(
            kind=SourceKind.UPLOADED_DOCUMENT,
            title=filename,
            retrieved_by=self.name,
            full_text_retrieved=True,
            retrieval_note="Plain-text file read verbatim.",
            abstract=text[:1000] or None,
        )
        for start, end, part in chunk(text):
            src.passages.append(
                Passage(source_id=src.id, text=part, char_start=start, char_end=end)
            )
        return src


class DocxParser(Provider):
    capability = CAPABILITY
    name = "docx"
    requires_credentials = False
    priority = 20
    extensions = (".docx",)

    def status(self) -> ProviderStatus:
        try:
            import docx  # noqa: F401
        except ImportError:
            return ProviderStatus(False, "python-docx is not installed.")
        return ProviderStatus(True, details={"extensions": list(self.extensions)})

    def parse(self, data: bytes, filename: str) -> Source:
        import docx

        try:
            document = docx.Document(io.BytesIO(data))
        except Exception as exc:  # noqa: BLE001
            raise UnsupportedContent(f"This file could not be read as a .docx: {exc}") from exc

        src = Source(
            kind=SourceKind.UPLOADED_DOCUMENT,
            title=(document.core_properties.title or filename),
            published=(
                document.core_properties.created.date()
                if document.core_properties.created
                else None
            ),
            retrieved_by=self.name,
            full_text_retrieved=True,
            retrieval_note="Word document text and tables extracted.",
        )
        if document.core_properties.author:
            from ...core.provenance import Author

            src.authors = [Author.parse(document.core_properties.author)]

        section = None
        buffer: list[str] = []

        def flush() -> None:
            nonlocal buffer
            text = normalise(" ".join(buffer))
            if text:
                for start, end, part in chunk(text):
                    src.passages.append(
                        Passage(source_id=src.id, text=part, section=section,
                                char_start=start, char_end=end)
                    )
            buffer = []

        for para in document.paragraphs:
            text = normalise(para.text)
            if not text:
                continue
            if para.style.name.startswith("Heading"):
                flush()
                section = text
            buffer.append(text)
        flush()

        tables = []
        for i, table in enumerate(document.tables[:30]):
            rows = [[normalise(c.text) for c in r.cells] for r in table.rows]
            if len(rows) >= 2:
                tables.append({"index": i, "header": rows[0], "rows": rows[1:]})
        src.extra["tables"] = tables
        src.abstract = normalise(" ".join(p.text for p in src.passages[:2]))[:1000] or None
        return src


class TabularParser(Provider):
    """CSV/TSV/Excel — parsed as *data*, and also indexed as readable text."""

    capability = CAPABILITY
    name = "tabular"
    requires_credentials = False
    priority = 15
    extensions = (".csv", ".tsv", ".xlsx", ".xls", ".xlsm")

    def status(self) -> ProviderStatus:
        try:
            import pandas  # noqa: F401
        except ImportError:
            return ProviderStatus(False, "pandas is not installed.")
        return ProviderStatus(True, details={"extensions": list(self.extensions)})

    def parse(self, data: bytes, filename: str) -> Source:
        import pandas as pd

        lower = filename.lower()
        try:
            if lower.endswith((".xlsx", ".xls", ".xlsm")):
                frames = pd.read_excel(io.BytesIO(data), sheet_name=None)
            else:
                sample = data[:8192].decode("utf-8", errors="replace")
                try:
                    sep = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
                except csv.Error:
                    sep = "\t" if lower.endswith(".tsv") else ","
                frames = {"data": pd.read_csv(io.BytesIO(data), sep=sep)}
        except Exception as exc:  # noqa: BLE001
            raise UnsupportedContent(f"This file could not be read as a table: {exc}") from exc

        src = Source(
            kind=SourceKind.DATASET,
            title=filename,
            retrieved_by=self.name,
            full_text_retrieved=True,
            retrieval_note="Tabular file loaded; values are used as-is, never imputed.",
        )
        sheets = {}
        for sheet, frame in frames.items():
            sheets[sheet] = {
                "rows": int(frame.shape[0]),
                "columns": [str(c) for c in frame.columns],
                "dtypes": {str(c): str(t) for c, t in frame.dtypes.items()},
                "missing": {str(c): int(frame[c].isna().sum()) for c in frame.columns},
            }
            preview = frame.head(30).to_csv(index=False)
            text = f"Sheet '{sheet}' — {frame.shape[0]} rows x {frame.shape[1]} columns.\n{preview}"
            for start, end, part in chunk(text, size=2000):
                src.passages.append(
                    Passage(source_id=src.id, text=part, section=str(sheet),
                            char_start=start, char_end=end)
                )
        src.extra["sheets"] = sheets
        src.abstract = "; ".join(
            f"{name}: {meta['rows']} rows x {len(meta['columns'])} cols" for name, meta in sheets.items()
        )
        return src
